#!/usr/bin/env python3
"""Validate private Excel data; optionally import it atomically into Supabase.

The workbook remains in memory. Logs contain counts and cell locations only.
No frontend module, public JSON, or credential file is generated.
"""

import argparse
import json
import os
from pathlib import Path
import re
import ssl
import sys
from types import SimpleNamespace
import unicodedata
from urllib.error import HTTPError, URLError
from urllib.parse import urlsplit
from urllib.request import HTTPRedirectHandler, HTTPSHandler, Request, build_opener
from zipfile import BadZipFile, ZipFile

try:
    import defusedxml  # noqa: F401: ensure hardened XML parsing is available
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Thiếu thư viện: cài scripts/requirements-import.txt trong virtualenv trước.", file=sys.stderr)
    raise SystemExit(2)


RECIPIENT_HEADERS = (
    "ID người nhận *", "Họ và tên *", "Mã sinh viên *", "Tên hiển thị",
    "Ảnh (URL hoặc đường dẫn)", "Giới thiệu", "Tính cách", "Kỷ niệm",
)
WISH_HEADERS = (
    "ID lời chúc *", "ID người nhận *", "Tiêu đề", "Người gửi", "Nội dung lời chúc *", "Thứ tự",
)
ID_PATTERN = re.compile(r"[A-Za-z0-9_-]{1,80}\Z", re.ASCII)
STUDENT_CODE_PATTERN = re.compile(r"[A-Z0-9._-]{3,64}\Z", re.ASCII)
JS_WHITESPACE = re.compile(r"[\t\n\v\f\r \u00a0\u1680\u2000-\u200a\u2028\u2029\u202f\u205f\u3000\ufeff]+")
REPOSITORY = Path(__file__).resolve().parents[2]
PRIVATE_DIRECTORY = Path(__file__).resolve().parents[1] / ".private"


class ImportProblem(Exception):
    """A safe error message: never include cell contents or provider bodies."""


def normalize_name(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", value)
    unaccented = "".join(character for character in decomposed if not "\u0300" <= character <= "\u036f")
    return JS_WHITESPACE.sub(" ", unaccented.replace("đ", "d").replace("Đ", "d").lower()).strip(" ")


def text_cell(cell, *, required: bool = False, maximum: int = 5000, default: str = "", single_line: bool = False) -> str:
    location = f"{cell.parent.title}!{cell.coordinate}"
    if cell.data_type in {"f", "e"}:
        raise ImportProblem(f"{location}: chỉ chấp nhận văn bản, không dùng công thức hoặc ô lỗi.")
    if cell.value is None:
        value = default
    elif not isinstance(cell.value, str):
        raise ImportProblem(f"{location}: hãy nhập dạng Văn bản (Text), không phải số/ngày tháng.")
    else:
        value = cell.value.strip()
        if not value:
            value = default
    if required and not value:
        raise ImportProblem(f"{location}: thiếu ô bắt buộc.")
    if len(value) > maximum:
        raise ImportProblem(f"{location}: vượt quá {maximum} ký tự.")
    if single_line and any(unicodedata.category(character) == "Cc" for character in value):
        raise ImportProblem(f"{location}: tên phải nằm trên một dòng và không có ký tự điều khiển.")
    if any(ord(character) < 32 and character not in "\n\r\t" for character in value):
        raise ImportProblem(f"{location}: có ký tự điều khiển không hợp lệ.")
    return value


def validate_id(cell) -> str:
    value = text_cell(cell, required=True, maximum=80)
    if not ID_PATTERN.fullmatch(value):
        raise ImportProblem(f"{cell.parent.title}!{cell.coordinate}: ID chỉ gồm chữ ASCII, số, gạch ngang hoặc gạch dưới.")
    return value


def worksheet_rows(sheet, headers: tuple[str, ...], maximum: int):
    if sheet.max_row > maximum + 4 or sheet.max_column != len(headers):
        raise ImportProblem(f"{sheet.title}: số dòng/cột vượt giới hạn hoặc mẫu đã bị đổi cấu trúc.")
    actual = tuple(cell.value for cell in next(sheet.iter_rows(min_row=4, max_row=4)))
    if actual != headers:
        raise ImportProblem(f"{sheet.title}: giữ nguyên tên và thứ tự các cột ở dòng 4 của mẫu.")
    for row_number, cells in enumerate(sheet.iter_rows(min_row=5), 5):
        if all(cell.value is None or cell.value == "" for cell in cells):
            continue
        # Read-only worksheets represent missing cells as unlocated EmptyCell.
        # Preserve locations for safe errors without echoing private contents.
        yield tuple(
            cell if hasattr(cell, "coordinate") else SimpleNamespace(
                value=None, data_type="n", parent=sheet, row=row_number,
                coordinate=f"{get_column_letter(column)}{row_number}",
            )
            for column, cell in enumerate(cells, 1)
        )


def validate_archive(path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        raise ImportProblem("Chỉ nhận file .xlsx; không nhận macro, CSV hoặc file đã đổi đuôi.")
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ImportProblem("File vượt giới hạn 20 MiB.")
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 80 * 1024 * 1024:
            raise ImportProblem("File Excel có cấu trúc quá lớn.")
        if any("vbaproject" in item.filename.lower() or item.filename.startswith("xl/externalLinks/") for item in entries):
            raise ImportProblem("Không nhận file có macro hoặc liên kết workbook ngoài.")


def validate_workbook(path: Path) -> tuple[list[dict], list[dict]]:
    validate_archive(path)
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if set(workbook.sheetnames) != {"Người nhận", "Lời chúc"} or len(workbook.sheetnames) != 2:
            raise ImportProblem("Workbook phải có đúng hai sheet: Người nhận và Lời chúc.")
        recipients = []
        recipient_ids = set()
        login_pairs = set()
        for cells in worksheet_rows(workbook["Người nhận"], RECIPIENT_HEADERS, 1000):
            identifier = validate_id(cells[0])
            if identifier in recipient_ids:
                raise ImportProblem(f"Người nhận!A{cells[0].row}: ID người nhận bị lặp.")
            full_name = text_cell(cells[1], required=True, maximum=120, single_line=True)
            normalized_name = normalize_name(full_name)
            if not normalized_name or len(normalized_name) > 200:
                raise ImportProblem(f"Người nhận!B{cells[1].row}: tên không hợp lệ sau chuẩn hóa.")
            student_code = text_cell(cells[2], required=True, maximum=64).upper()
            if not STUDENT_CODE_PATTERN.fullmatch(student_code):
                raise ImportProblem(f"Người nhận!C{cells[2].row}: mã phải có 3–64 ký tự ASCII A–Z, 0–9, chấm, gạch dưới hoặc gạch ngang.")
            login_pair = (normalized_name, student_code)
            if login_pair in login_pairs:
                raise ImportProblem(f"Người nhận!C{cells[2].row}: tên và mã đăng nhập bị lặp.")
            image = text_cell(cells[4], maximum=500)
            if image and (any(ord(character) < 32 for character in image) or not (
                (image.startswith("/") and not image.startswith("//")) or
                (urlsplit(image).scheme == "https" and bool(urlsplit(image).hostname) and not urlsplit(image).username)
            )):
                raise ImportProblem(f"Người nhận!E{cells[4].row}: ảnh phải là đường dẫn /… hoặc URL HTTPS hợp lệ.")
            recipients.append({
                "id": identifier,
                "full_name": full_name,
                "normalized_name": normalized_name,
                "student_code": student_code,
                "display_name": text_cell(cells[3], maximum=120, default=full_name, single_line=True),
                "image": image,
                "introduction": text_cell(cells[5]),
                "personality": text_cell(cells[6]),
                "memory": text_cell(cells[7]),
            })
            recipient_ids.add(identifier)
            login_pairs.add(login_pair)
        wishes = []
        wish_ids = set()
        for cells in worksheet_rows(workbook["Lời chúc"], WISH_HEADERS, 10000):
            identifier = validate_id(cells[0])
            if identifier in wish_ids:
                raise ImportProblem(f"Lời chúc!A{cells[0].row}: ID lời chúc bị lặp.")
            recipient_id = validate_id(cells[1])
            if recipient_id not in recipient_ids:
                raise ImportProblem(f"Lời chúc!B{cells[1].row}: chưa có ID người nhận tương ứng trong sheet Người nhận.")
            order = cells[5].value
            if order is None or order == "":
                order = cells[5].row - 5
            if isinstance(order, str) and re.fullmatch(r"[0-9]{1,10}", order.strip(), flags=re.ASCII):
                order = int(order.strip())
            if isinstance(order, float) and order.is_integer():
                order = int(order)
            if type(order) is not int or not 0 <= order <= 2147483647 or cells[5].data_type in {"f", "e"}:
                raise ImportProblem(f"Lời chúc!F{cells[5].row}: thứ tự phải là số nguyên từ 0 đến 2147483647.")
            wishes.append({
                "id": identifier,
                "recipient_id": recipient_id,
                "title": text_cell(cells[2], maximum=200, default="Lời chúc dành cho cậu"),
                "sender_name": text_cell(cells[3], maximum=120, default="Một người bạn DM66D"),
                "message": text_cell(cells[4], required=True, maximum=10000),
                "sort_order": order,
            })
            wish_ids.add(identifier)
        return recipients, wishes
    finally:
        workbook.close()


class NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, request, file_pointer, code, message, headers, new_url):
        # Never forward API credentials or a private workbook to another origin.
        return None


def apply_import(recipients: list[dict], wishes: list[dict], environment=None, opener=None) -> None:
    environment = os.environ if environment is None else environment
    url = environment.get("SUPABASE_URL", "").strip()
    key = environment.get("SUPABASE_SERVICE_ROLE_KEY", "").strip()
    try:
        parsed = urlsplit(url)
        valid = (parsed.scheme == "https" and parsed.hostname and
                 parsed.hostname.endswith(".supabase.co") and parsed.port in {None, 443} and
                 not parsed.username and not parsed.password and parsed.path in {"", "/"} and
                 not parsed.query and not parsed.fragment)
    except ValueError:
        valid = False
    if not valid or not key:
        raise ImportProblem("Cần SUPABASE_URL HTTPS của project Supabase và SUPABASE_SERVICE_ROLE_KEY trong biến môi trường máy chủ.")
    key_format = environment.get("SUPABASE_KEY_FORMAT", "auto").strip().lower()
    if key_format not in {"auto", "legacy", "modern"}:
        raise ImportProblem("SUPABASE_KEY_FORMAT phải là auto, legacy hoặc modern.")
    if not recipients:
        raise ImportProblem("Mẫu đang trống; chưa nhập dữ liệu nào lên máy chủ.")
    body = json.dumps({"p_recipients": recipients, "p_wishes": wishes}, ensure_ascii=False).encode("utf-8")
    request = Request(url.rstrip("/") + "/rest/v1/rpc/dm66_import_workbook", data=body, method="POST")
    request.add_header("apikey", key)
    request.add_header("Content-Type", "application/json; charset=utf-8")
    request.add_header("Accept", "application/json")
    if key_format == "legacy" or (key_format == "auto" and len(key.split(".")) == 3):
        request.add_header("Authorization", "Bearer " + key)
    if opener is None:
        # create_default_context preserves OS/SSL_CERT_FILE CA verification;
        # build_opener also preserves the machine's configured HTTPS proxy.
        opener = build_opener(NoRedirect(), HTTPSHandler(context=ssl.create_default_context()))
    try:
        with opener.open(request, timeout=60) as response:
            result = json.loads(response.read(4096))
    except HTTPError as error:
        raise ImportProblem(f"Supabase từ chối giao dịch (HTTP {error.code}); kiểm tra migration, quyền service_role và cấu hình Preview. Không ghi một phần dữ liệu.") from None
    except (URLError, TimeoutError, OSError):
        raise ImportProblem("Không xác nhận được giao dịch do lỗi kết nối/TLS. Giữ xác minh chứng chỉ; kiểm tra kết nối rồi nhập lại cùng các ID.") from None
    except (ValueError, UnicodeError):
        raise ImportProblem("Phản hồi nhập dữ liệu không hợp lệ; kiểm tra migration. Có thể nhập lại an toàn với cùng các ID.") from None
    if result != {"recipients": len(recipients), "wishes": len(wishes)}:
        raise ImportProblem("Máy chủ chưa xác nhận đúng số lượng; kiểm tra migration. Có thể nhập lại an toàn với cùng các ID.")


def ensure_private_location(path: Path) -> None:
    if path.is_relative_to(REPOSITORY) and not path.is_relative_to(PRIVATE_DIRECTORY):
        raise ImportProblem("File đã điền trong repository phải nằm trong dm66-website/.private/ (được gitignore); nên lưu ngoài repository. Không đặt trong public/ hoặc src/.")


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Kiểm tra Excel riêng tư; mặc định chỉ kiểm tra, không ghi mạng.")
    parser.add_argument("workbook", type=Path, help="File .xlsx đã điền, ngoài repository hoặc trong .private/")
    parser.add_argument("--apply", action="store_true", help="Nhập toàn bộ workbook trong một giao dịch Supabase bằng biến môi trường máy chủ")
    arguments = parser.parse_args(argv)
    try:
        path = arguments.workbook.resolve(strict=True)
        ensure_private_location(path)
        recipients, wishes = validate_workbook(path)
        print(f"Hợp lệ: {len(recipients)} người nhận, {len(wishes)} lời chúc.")
        if arguments.apply:
            apply_import(recipients, wishes)
            print(f"Đã nhập nguyên tử: {len(recipients)} người nhận, {len(wishes)} lời chúc. Các ID không có trong file vẫn được giữ nguyên.")
        else:
            print("Chỉ kiểm tra; chưa gửi dữ liệu hoặc ghi lên máy chủ. Thêm --apply để nhập vào Supabase Preview đã cấu hình.")
        return 0
    except ImportProblem as error:
        print(f"Không nhập dữ liệu: {error}", file=sys.stderr)
    except (OSError, BadZipFile, ValueError, KeyError):
        print("Không nhập dữ liệu: không đọc được file Excel; kiểm tra quyền truy cập và định dạng .xlsx đúng mẫu.", file=sys.stderr)
    except Exception:
        # Untrusted Excel/XML and provider failures must never dump a traceback
        # with private filenames, cell contents, headers, or response details.
        print("Không nhập dữ liệu: không xử lý được file Excel hoặc phản hồi máy chủ; kiểm tra định dạng mẫu và cấu hình Preview.", file=sys.stderr)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
