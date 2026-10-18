#!/usr/bin/env python3
"""Import the class's two-photo memories workbook and original JPEGs offline.

Only Du_lieu_ky_niem is data. The other workbook sheets are never interpreted.
No source workbook or ZIP is copied into the public site.
"""

import argparse
from datetime import date, datetime
import json
from pathlib import Path, PurePosixPath
import re
import stat
from zipfile import ZipFile

import defusedxml  # noqa: F401: enable openpyxl's hardened XML parser
from openpyxl import load_workbook


HEADERS = (
    "STT (thứ tự)", "id *", "title – Tên kỷ niệm *", "date – Thời gian / dịp",
    "image – Tên file ảnh *", "description – Mô tả kỷ niệm *",
)
PLACEHOLDER = "/images/memories-placeholder.svg"
ALIASES = {"cnxkh1.jpg": "cnxhkh1.jpg"}
SAFE_FILENAME = re.compile(r"[a-z0-9_-]+\.jpe?g\Z", re.ASCII)
SAFE_ID = re.compile(r"[a-zA-Z0-9_-]{1,80}\Z", re.ASCII)


class ImportProblem(ValueError):
    pass


def audit_archive(archive, maximum_size=100 * 1024 * 1024):
    """Reject traversal, links, encryption, excessive expansion and file counts."""
    entries = archive.infolist()
    if len(entries) > 2000 or sum(entry.file_size for entry in entries) > maximum_size:
        raise ImportProblem("Archive exceeds the supported size or file count.")
    for entry in entries:
        path = PurePosixPath(entry.filename)
        if (path.is_absolute() or ".." in path.parts or "\\" in entry.filename
                or ":" in entry.filename or "\x00" in entry.filename):
            raise ImportProblem("Archive contains an unsafe path.")
        if stat.S_ISLNK(entry.external_attr >> 16) or entry.flag_bits & 1:
            raise ImportProblem("Archive links and encrypted files are unsupported.")
        if entry.file_size > 30 * 1024 * 1024:
            raise ImportProblem("An archive entry exceeds the supported size.")


def jpeg_assets(zip_path):
    assets = {}
    with ZipFile(zip_path) as archive:
        audit_archive(archive)
        for entry in archive.infolist():
            if entry.is_dir():
                continue
            name = PurePosixPath(entry.filename).name.lower()
            if not SAFE_FILENAME.fullmatch(name):
                raise ImportProblem("Photo archive must contain JPEGs only.")
            if name in assets:
                raise ImportProblem("Photo archive has duplicate filenames ignoring case.")
            content = archive.read(entry)
            if not content.startswith(b"\xff\xd8\xff") or not content.endswith(b"\xff\xd9"):
                raise ImportProblem(f"{name}: file is not an intact JPEG.")
            assets[name] = content
    return assets


def cell_text(cell):
    if cell.data_type in {"f", "e"}:
        raise ImportProblem(f"{cell.coordinate}: formulas and cell errors are unsupported.")
    if cell.value is None:
        return ""
    if not isinstance(cell.value, str):
        raise ImportProblem(f"{cell.coordinate}: expected text.")
    value = cell.value.strip()
    if len(value) > 20000 or any(ord(char) < 32 and char not in "\n\r\t" for char in value):
        raise ImportProblem(f"{cell.coordinate}: invalid or excessively long text.")
    return value


def build_data(workbook_path, zip_path):
    assets = jpeg_assets(zip_path)
    with ZipFile(workbook_path) as archive:
        audit_archive(archive, maximum_size=50 * 1024 * 1024)
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    try:
        if "Du_lieu_ky_niem" not in workbook:
            raise ImportProblem("Missing Du_lieu_ky_niem worksheet.")
        sheet = workbook["Du_lieu_ky_niem"]
        # The supplied workbook has no dimension metadata; styled empty rows
        # must not be mistaken for actual memories or stop iteration early.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(max_col=6)
        if tuple(cell.value for cell in next(iterator)) != HEADERS:
            raise ImportProblem("The memories header does not match the template.")
        records, ids, orders = [], set(), set()
        missing, renamed, title_fallbacks, used = [], [], [], set()
        for row_number, cells in enumerate(iterator, 2):
            if row_number > 2001:
                raise ImportProblem("Too many worksheet rows.")
            if all(cell.value in (None, "") for cell in cells):
                continue
            order = cells[0].value
            if (cells[0].data_type in {"f", "e"} or isinstance(order, bool)
                    or not isinstance(order, (int, float)) or not 1 <= order <= 10000
                    or int(order) != order or order in orders):
                raise ImportProblem(f"A{row_number}: expected a unique positive integer.")
            identifier = cell_text(cells[1])
            if not SAFE_ID.fullmatch(identifier) or identifier in ids:
                raise ImportProblem(f"B{row_number}: invalid or duplicate memory ID.")
            title, description = cell_text(cells[2]), cell_text(cells[5])
            if not title:
                if identifier == "ky-niem-13" and "Boy Day 2026" in description:
                    title = "Boy Day 2026"
                    title_fallbacks.append(identifier)
                else:
                    raise ImportProblem(f"C{row_number}: missing event title.")
            image_names = [line.strip().lower() for line in cell_text(cells[4]).splitlines() if line.strip()]
            if len(image_names) != 2 or any(not SAFE_FILENAME.fullmatch(name) for name in image_names):
                raise ImportProblem(f"E{row_number}: expected two JPEG filenames on separate lines.")
            images = []
            for index, original_name in enumerate(image_names, 1):
                name = ALIASES.get(original_name, original_name)
                if name in assets:
                    images.append({"src": f"/images/memories/{name}", "altText": f"{title} — ảnh kỷ niệm {index}"})
                    used.add(name)
                    if name != original_name:
                        renamed.append({"id": identifier, "from": original_name, "to": name})
                else:
                    images.append({"src": PLACEHOLDER, "altText": f"{title} — đang chờ ảnh kỷ niệm {index}", "isPlaceholder": True})
                    missing.append({"id": identifier, "file": original_name})
            record = {"id": identifier, "title": title, "images": images, "description": description}
            date_cell = cells[3]
            if isinstance(date_cell.value, (date, datetime)):
                record["date"] = date_cell.value.strftime("%d/%m/%Y")
            else:
                date_text = cell_text(date_cell)
                if date_text:
                    record["date"] = date_text
            records.append((order, record))
            ids.add(identifier)
            orders.add(order)
        if not records:
            raise ImportProblem("No memories found.")
        records.sort(key=lambda pair: pair[0])
        memories = [record for _, record in records]
        return memories, {name: assets[name] for name in sorted(used)}, {
            "missing": missing, "aliases": renamed, "titleFallbacks": title_fallbacks,
            "unusedImages": sorted(set(assets) - used),
            "blankDescriptions": [memory["id"] for memory in memories if not memory["description"]],
        }
    finally:
        workbook.close()


def typescript(memories):
    return '''export interface MemoryImage {
  src: string;
  altText: string;
  isPlaceholder?: boolean;
}

export interface Memory {
  id: string;
  title: string;
  date?: string;
  images: MemoryImage[];
  description: string;
}

// Imported from the class workbook; keep the original descriptions intact.
export const memories: Memory[] = ''' + json.dumps(memories, ensure_ascii=False, indent=2) + ";\n"


def report_markdown(memories, assets, diagnostics, workbook_name, zip_name):
    lines = [
        "# Nhập dữ liệu kỷ niệm lớp", "",
        f"Nguồn: `{workbook_name}` và `{zip_name}` do người dùng cung cấp.", "",
        f"- Đã nhập **{len(memories)} sự kiện**, mỗi sự kiện có hai vị trí ảnh.",
        f"- Đã sử dụng **{len(assets)} ảnh JPEG gốc**, giữ nguyên từng byte; chỉ chuẩn hóa tên file về chữ thường.",
        f"- Có **{len(diagnostics['missing'])} ảnh thiếu** dùng hình khung trống trung tính, không lấy ảnh sự kiện khác thay vào.",
        "- Chỉ đọc sheet `Du_lieu_ky_niem`; không nhập ví dụ hoặc diễn giải ghi chú thành dữ liệu.",
        "- Giữ đầy đủ tiêu đề và mô tả nguyên văn, chỉ bỏ khoảng trắng ở hai đầu ô.",
        "- Không sao chép workbook hoặc ZIP gốc vào thư mục công khai.", "",
        "## Các điểm đã xử lý", "",
    ]
    for alias in diagnostics["aliases"]:
        lines.append(f"- `{alias['id']}`: tên ảnh `{alias['from']}` trong Excel được đối chiếu rõ ràng với `{alias['to']}` trong ZIP.")
    for identifier in diagnostics["titleFallbacks"]:
        lines.append(f"- `{identifier}`: ô tên sự kiện trống; dùng **Boy Day 2026** vì mô tả ghi rõ tên này.")
    for identifier in diagnostics["blankDescriptions"]:
        lines.append(f"- `{identifier}`: mô tả nguồn trống, giữ trống và không tự viết thêm câu chuyện.")
    for missing in diagnostics["missing"]:
        lines.append(f"- `{missing['id']}`: ZIP không có `{missing['file']}`; giữ sự kiện và hiển thị trạng thái đang chờ ảnh.")
    if all("date" not in memory for memory in memories):
        lines.append("- Toàn bộ ô thời gian trong Excel trống; không tự gán ngày hoặc năm.")
    lines.extend(["", "## Danh sách sự kiện", "", "| ID | Tên đầy đủ | Ảnh thật |", "| --- | --- | ---: |"])
    for memory in memories:
        title = memory["title"].replace("|", "\\|").replace("\n", "<br>")
        count = sum(not image.get("isPlaceholder", False) for image in memory["images"])
        lines.append(f"| `{memory['id']}` | {title} | {count}/2 |")
    lines.extend(["", "## Chạy lại", "", "Từ `dm66-website`, dùng môi trường Python đã cài `scripts/requirements-import.txt`:", "", "```bash", f".venv-import/bin/python scripts/import_memories.py '/duong-dan/{workbook_name}' '/duong-dan/{zip_name}'", "```", "", "Công cụ kiểm tra ZIP, từ chối đường dẫn không an toàn, symlink, công thức trong các ô dữ liệu và ảnh bị trùng tên; không giải nén tùy ý. Những ảnh hiện có có nội dung khác sẽ không bị ghi đè.", ""])
    return "\n".join(lines)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("photos", type=Path)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args(argv)
    memories, assets, diagnostics = build_data(args.workbook, args.photos)
    directory = args.root / "public/images/memories"
    data_file = args.root / "src/data/memories.ts"
    # Validate all inputs and preserve pre-existing differing files before writing.
    for name, content in assets.items():
        path = directory / name
        if path.exists() and path.read_bytes() != content:
            raise ImportProblem(f"Refusing to overwrite a different existing photo: {name}.")
    directory.mkdir(parents=True, exist_ok=True)
    for name, content in assets.items():
        (directory / name).write_bytes(content)
    data_file.parent.mkdir(parents=True, exist_ok=True)
    data_file.write_text(typescript(memories), encoding="utf-8")
    report_file = args.root / "docs/memories-import-report.md"
    report_file.parent.mkdir(parents=True, exist_ok=True)
    report_file.write_text(report_markdown(memories, assets, diagnostics, args.workbook.name, args.photos.name), encoding="utf-8")
    print(json.dumps({"memories": len(memories), "images": len(assets), **diagnostics}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
