#!/usr/bin/env python3
"""Create the public, empty Excel template. Never use this for private data."""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from import_private_wishes import RECIPIENT_HEADERS, WISH_HEADERS


def create_template(destination: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "DM66D"
    workbook.properties.title = "DM66D · Mẫu người nhận và lời chúc riêng"
    workbook.properties.created = datetime(2026, 10, 18)
    workbook.properties.modified = datetime(2026, 10, 18)
    sheets = [
        (
            "Người nhận",
            RECIPIENT_HEADERS,
            [24, 30, 24, 26, 42, 50, 50, 50],
            "Bắt đầu từ dòng 5 · * bắt buộc · Mã sinh viên phải là văn bản để giữ số 0 đầu · Giữ ID ổn định khi nhập lại.",
            [
                "ID duy nhất, 1–80 ký tự chữ ASCII, số, dấu gạch ngang hoặc gạch dưới. Không dùng mã sinh viên làm ID.",
                "Họ và tên dùng khi đăng nhập; được so sánh không phân biệt hoa/thường, dấu và khoảng trắng lặp.",
                "Nhập dạng Văn bản (Text), 3–64 ký tự A–Z, 0–9, dấu chấm, gạch dưới hoặc gạch ngang. Không công thức. Không công khai file đã điền.",
                "Tên hiển thị sau đăng nhập; để trống sẽ dùng họ và tên.",
                "Ảnh: đường dẫn /images/… hoặc URL https://…; để trống nếu chưa có. Chỉ sử dụng ảnh được phép chia sẻ.",
                "Giới thiệu tùy chọn, chỉ trả về cho đúng người sau đăng nhập.",
                "Tính cách tùy chọn.",
                "Kỷ niệm tùy chọn.",
            ],
        ),
        (
            "Lời chúc",
            WISH_HEADERS,
            [24, 24, 36, 28, 80, 18],
            "Bắt đầu từ dòng 5 · * bắt buộc · Nhiều dòng có thể dùng cùng ID người nhận · Mỗi lời chúc có ID riêng, ổn định.",
            [
                "ID duy nhất của lời chúc; giữ nguyên khi sửa để cập nhật thay vì tạo bản sao.",
                "Phải trùng ID ở sheet Người nhận; một người nhận có thể có nhiều lời chúc.",
                "Tiêu đề lời chúc; để trống sẽ dùng Lời chúc dành cho cậu.",
                "Tên người gửi (tùy chọn); để trống để giữ lời chúc ẩn danh.",
                "Nội dung bắt buộc; có thể xuống dòng bằng Alt+Enter. Nhập văn bản, không công thức.",
                "Số nguyên từ 0 trở lên; nhỏ hơn hiển thị trước. Để trống sẽ dùng thứ tự dòng trong sheet.",
            ],
        ),
    ]
    for title, headers, widths, instruction, comments in sheets:
        sheet = workbook.create_sheet(title)
        last_column = len(headers)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
        sheet.cell(1, 1, f"DM66D · {title.upper()}")
        sheet.cell(1, 1).font = Font(name="Calibri", size=20, bold=True, color="6F283C")
        sheet.row_dimensions[1].height = 34
        for row, note in [
            (2, instruction),
            (3, "BẢO MẬT: file đã điền chỉ lưu ngoài repository hoặc trong .private/; không tải lên frontend, GitHub hay gửi qua chat."),
        ]:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
            cell = sheet.cell(row, 1, note)
            cell.font = Font(name="Calibri", size=11, color="6F283C" if row == 3 else "52464A", italic=row == 3)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row].height = 32
        for column, (header, width, comment) in enumerate(zip(headers, widths, comments), 1):
            cell = sheet.cell(4, column, header)
            cell.font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
            cell.fill = PatternFill("solid", fgColor="8D3B53")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.comment = Comment(comment, "DM66D")
            sheet.column_dimensions[cell.column_letter].width = width
        sheet.row_dimensions[4].height = 36
        border = Border(bottom=Side(style="hair", color="E8D8DD"))
        for row in range(5, 55):
            sheet.row_dimensions[row].height = 30
            for column in range(1, last_column + 1):
                cell = sheet.cell(row, column)
                cell.number_format = "@"
                cell.font = Font(name="Calibri", size=11, color="3E3438")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border
                cell.fill = PatternFill("solid", fgColor="FFF5F7" if row % 2 else "FFFFFF")
        sheet.freeze_panes = "C5"
        sheet.auto_filter.ref = f"A4:{sheet.cell(54, last_column).column_letter}54"
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = "1:4"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        if title == "Lời chúc":
            validation = DataValidation(type="whole", operator="between", formula1=0, formula2=2147483647, allow_blank=True)
            validation.errorTitle = "Thứ tự không hợp lệ"
            validation.error = "Nhập số nguyên từ 0 trở lên."
            validation.showErrorMessage = True
            sheet.add_data_validation(validation)
            validation.add("F5:F10004")
            for row in range(5, 55):
                sheet.cell(row, 6).number_format = "0"
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)


if __name__ == "__main__":
    target = Path(__file__).resolve().parents[1] / "public/downloads/DM66D-private-wishes-template.xlsx"
    create_template(target)
    print("Đã tạo Excel mẫu trống gồm 2 sheet: Người nhận, Lời chúc.")
