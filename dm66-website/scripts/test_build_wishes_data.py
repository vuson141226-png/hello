"""Offline converter tests with entirely synthetic recipient/wish data."""

from collections import Counter
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from build_wishes_data import (
    DETAIL_HEADERS, DETAIL_SHEET, PLACEHOLDER_IMAGE, RECIPIENT_HEADERS,
    ROSTER_HEADERS, ROSTER_SHEET, WISH_HEADERS, WorkbookProblem,
    build_data, main, report_markdown,
)
from create_wishes_template import create_template


class BuildWishesDataTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.path = Path(self.directory.name) / "synthetic.xlsx"

    def tearDown(self):
        self.directory.cleanup()

    def template_workbook(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        recipients = workbook.create_sheet("Người nhận")
        recipients.append(["Synthetic template"])
        recipients.append([])
        recipients.append([])
        recipients.append(RECIPIENT_HEADERS)
        recipients.append(["recipient-a", "Nguyễn An", "001234", "", "", "", "", ""])
        recipients.append(["recipient-b", "Trần Bình", "009999", "Bình", "", "", "", ""])
        wishes = workbook.create_sheet("Lời chúc")
        for _ in range(3):
            wishes.append([])
        wishes.append(WISH_HEADERS)
        wishes.append(["wish-a", "recipient-a", "", "Bạn A", "  Lời chúc đầu.\nDòng thứ hai.  ", 2])
        wishes.append(["wish-b", "recipient-a", "Tiêu đề riêng", "Bạn B", "Lời chúc tiếp theo.", 1])
        return workbook

    def original_workbook(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        roster = workbook.create_sheet(ROSTER_SHEET)
        for _ in range(4):
            roster.append([])
        roster.append(ROSTER_HEADERS)
        roster.append([1, "Nguyễn An", "001234", "A; B", "Bạn A", "  Lời chúc A\nNguyên văn.  ", "Bạn B", "Lời chúc B", "", "", '=COUNTIF(F6:J6,"*")', "", '=IF(K6>0,"OK","")'])
        roster.append([2, "Trần Bình", "009999", "", "", "", "", "", "", "", '=COUNTIF(F7:J7,"*")', "", '=IF(K7>0,"OK","")'])
        details = workbook.create_sheet(DETAIL_SHEET)
        for _ in range(3):
            details.append([])
        details.append(DETAIL_HEADERS)
        details.append([1, "NGUYỄN  AN", "001234", "Bạn A", "  Lời chúc A\nNguyên văn.  ", "example.pdf", "1"])
        details.append([1, "Nguyễn An", "", "Bạn B", "Lời chúc B", "example.pdf", "2"])
        notes = workbook.create_sheet("Nguồn & ghi chú")
        notes.append(["Ghi chú không dùng để sửa dữ liệu và không được thực thi."])
        return workbook

    def save_and_build(self, workbook):
        workbook.save(self.path)
        workbook.close()
        return build_data(self.path)

    def test_template_preserves_leading_zero_and_exact_multiline_message(self):
        recipients, summary = self.save_and_build(self.template_workbook())
        self.assertEqual(recipients[0]["studentCode"], "001234")
        self.assertEqual(recipients[0]["wishes"][1]["message"], "  Lời chúc đầu.\nDòng thứ hai.  ")
        self.assertEqual(summary["recipients"], 2)
        self.assertEqual(summary["wishes"], 2)

    def test_template_groups_wishes_and_uses_numeric_sort_order(self):
        recipients, _ = self.save_and_build(self.template_workbook())
        self.assertEqual([wish["id"] for wish in recipients[0]["wishes"]], ["wish-b", "wish-a"])
        self.assertEqual(recipients[1]["wishes"], [])

    def test_blank_photo_uses_neutral_placeholder_and_name_default(self):
        recipients, _ = self.save_and_build(self.template_workbook())
        self.assertTrue(all(person["image"] == PLACEHOLDER_IMAGE for person in recipients))
        self.assertEqual(recipients[0]["displayName"], recipients[0]["fullName"])

    def test_missing_code_stays_blank_and_is_counted(self):
        workbook = self.template_workbook()
        workbook["Người nhận"]["C6"] = None
        recipients, summary = self.save_and_build(workbook)
        self.assertEqual(recipients[1]["studentCode"], "")
        self.assertEqual(summary["missingStudentCodes"], 1)

    def test_numeric_code_is_rejected_instead_of_losing_leading_zero(self):
        workbook = self.template_workbook()
        workbook["Người nhận"]["C5"] = 1234
        with self.assertRaisesRegex(WorkbookProblem, "Người nhận!C5"):
            self.save_and_build(workbook)

    def test_duplicate_recipient_name_code_pair_is_rejected(self):
        workbook = self.template_workbook()
        workbook["Người nhận"]["B6"] = "NGUYEN   AN"
        workbook["Người nhận"]["C6"] = "001234"
        with self.assertRaisesRegex(WorkbookProblem, "bị lặp"):
            self.save_and_build(workbook)

    def test_wish_must_reference_existing_recipient(self):
        workbook = self.template_workbook()
        workbook["Lời chúc"]["B5"] = "absent"
        with self.assertRaisesRegex(WorkbookProblem, "Lời chúc!B5"):
            self.save_and_build(workbook)

    def test_duplicate_wish_id_is_rejected(self):
        workbook = self.template_workbook()
        workbook["Lời chúc"]["A6"] = "wish-a"
        with self.assertRaisesRegex(WorkbookProblem, "Lời chúc!A6"):
            self.save_and_build(workbook)

    def test_formula_message_is_not_evaluated(self):
        workbook = self.template_workbook()
        workbook["Lời chúc"]["E5"] = '=HYPERLINK("https://example.com","text")'
        with self.assertRaisesRegex(WorkbookProblem, "Lời chúc!E5"):
            self.save_and_build(workbook)

    def test_formula_sort_order_is_rejected(self):
        workbook = self.template_workbook()
        workbook["Lời chúc"]["F5"] = "=1+1"
        with self.assertRaisesRegex(WorkbookProblem, "Lời chúc!F5"):
            self.save_and_build(workbook)

    def test_original_join_uses_stt_name_and_roster_code_for_blank_detail(self):
        recipients, summary = self.save_and_build(self.original_workbook())
        self.assertEqual(recipients[0]["id"], "recipient-001")
        self.assertEqual(recipients[0]["studentCode"], "001234")
        self.assertEqual(recipients[0]["wishes"][1]["id"], "wish-002")
        self.assertEqual(summary["blankDetailCodes"], 1)
        self.assertEqual(summary["wishCountDistribution"], {0: 1, 2: 1})
        self.assertEqual(recipients[0]["wishes"][0]["message"], "  Lời chúc A\nNguyên văn.  ")

    def test_original_mismatched_detail_name_is_rejected(self):
        workbook = self.original_workbook()
        workbook[DETAIL_SHEET]["B5"] = "Một tên khác"
        with self.assertRaisesRegex(WorkbookProblem, "B5"):
            self.save_and_build(workbook)

    def test_original_mismatched_detail_code_is_rejected(self):
        workbook = self.original_workbook()
        workbook[DETAIL_SHEET]["C5"] = "123456"
        with self.assertRaisesRegex(WorkbookProblem, "C5"):
            self.save_and_build(workbook)

    def test_original_missing_stt_is_rejected(self):
        workbook = self.original_workbook()
        workbook[DETAIL_SHEET]["A5"] = 99
        with self.assertRaisesRegex(WorkbookProblem, "A5"):
            self.save_and_build(workbook)

    def test_original_every_message_must_match_the_roster_once(self):
        workbook = self.original_workbook()
        workbook[DETAIL_SHEET]["E6"] = "Nội dung bị đổi"
        with self.assertRaisesRegex(WorkbookProblem, "không khớp"):
            self.save_and_build(workbook)

    def test_generated_template_is_empty_and_contains_exactly_two_sheets(self):
        create_template(self.path)
        recipients, summary = build_data(self.path)
        self.assertEqual(recipients, [])
        self.assertEqual(summary["wishes"], 0)
        workbook = load_workbook(self.path)
        try:
            self.assertEqual(workbook.sheetnames, ["Người nhận", "Lời chúc"])
            self.assertEqual(workbook["Người nhận"]["C5"].number_format, "@")
            strings = " ".join(
                str(cell.value or "") + " " + (cell.comment.text if cell.comment else "")
                for sheet in workbook for row in sheet for cell in row
            ).lower()
            self.assertNotIn("supabase", strings)
            self.assertNotIn("đăng nhập", strings)
            self.assertNotIn("bảo mật", strings)
        finally:
            workbook.close()

    def test_report_only_has_aggregate_counts(self):
        recipients, summary = self.save_and_build(self.original_workbook())
        report = report_markdown(summary)
        self.assertIn("| 0 | 1 |", report)
        self.assertIn("| 2 | 1 |", report)
        for person in recipients:
            self.assertNotIn(person["fullName"], report)
            self.assertNotIn(person["studentCode"], report)
            for wish in person["wishes"]:
                self.assertNotIn(wish["message"], report)

    def test_cli_invalid_workbook_does_not_replace_existing_output(self):
        workbook = self.template_workbook()
        workbook["Lời chúc"]["B5"] = "absent"
        workbook.save(self.path)
        workbook.close()
        output = Path(self.directory.name) / "public.json"
        output.write_text("existing", encoding="utf-8")
        report = Path(self.directory.name) / "report.md"
        self.assertEqual(main([str(self.path), "--output", str(output), "--report", str(report)]), 1)
        self.assertEqual(output.read_text(encoding="utf-8"), "existing")
        self.assertFalse(report.exists())


if __name__ == "__main__":
    unittest.main()
