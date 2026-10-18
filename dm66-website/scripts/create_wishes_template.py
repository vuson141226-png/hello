#!/usr/bin/env python3
"""Create an empty, reusable two-sheet workbook for the website."""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from build_wishes_data import RECIPIENT_HEADERS, WISH_HEADERS


def create_template(destination: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "DM66D"
    workbook.properties.title = "DM66D · Mẫu người nhận và lời chúc"
    workbook.properties.created = datetime(2026, 10, 18)
    workbook.properties.modified = datetime(2026, 10, 18)
    sheets = [
        (
            "Người nhận", RECIPIENT_HEADERS, [24, 30, 24, 26, 42, 50, 50, 50],
            "Bắt đầu từ dòng 5 · * bắt buộc · Mã sinh viên nhập dạng Văn bản (Text) để giữ số 0 đầu.",
            "Mỗi người nhận một dòng. Tra cứu bằng họ và tên + mã sinh viên.",
            [
                "ID duy nhất, gồm chữ ASCII, số, gạch ngang hoặc gạch dưới; ví dụ nguoi-nhan-001.",
                "Họ và tên dùng để đối chiếu. Có thể nhập không dấu khi tra cứu trên website.",
                "Nhập dạng Văn bản (Text) để giữ nguyên mã, kể cả số 0 đầu. Không dùng công thức.",
                "Để trống sẽ hiển thị họ và tên.",
                "Tùy chọn: đường dẫn /images/… hoặc URL https://…; để trống dùng ảnh minh họa chung.",
                "Giới thiệu tùy chọn.", "Tính cách tùy chọn.", "Kỷ niệm tùy chọn.",
            ],
        ),
        (
            "Lời chúc", WISH_HEADERS, [24, 24, 36, 28, 80, 18],
            "Bắt đầu từ dòng 5 · * bắt buộc · Mỗi lời chúc một dòng, mỗi lời chúc có ID riêng.",
            "Một người có nhiều lời chúc: lặp lại cùng ID người nhận ở nhiều dòng.",
            [
                "ID duy nhất của lời chúc; ví dụ loi-chuc-001. Giữ ID khi cập nhật nội dung.",
                "Phải trùng ID ở sheet Người nhận. Có thể dùng cùng một ID người nhận cho nhiều lời chúc.",
                "Tiêu đề tùy chọn; để trống sẽ tạo tiêu đề theo người gửi.",
                "Tên người gửi; để trống hiển thị Một người bạn DM66D.",
                "Nội dung nguyên văn. Xuống dòng bằng Alt+Enter; không nhập công thức.",
                "Số nguyên không âm; số nhỏ hiển thị trước. Để trống sẽ dùng thứ tự dòng.",
            ],
        ),
    ]
    for title, headers, widths, instruction, hint, comments in sheets:
        sheet = workbook.create_sheet(title)
        last_column = len(headers)
        for row, note in [(1, f"DM66D · {title.upper()}"), (2, instruction), (3, hint)]:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
            cell = sheet.cell(row, 1, note)
            cell.font = Font(name="Calibri", size=20 if row == 1 else 11, bold=row == 1, color="6F283C")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row].height = 34 if row == 1 else 32
        for column, (header, width, comment) in enumerate(zip(headers, widths, comments), 1):
            cell = sheet.cell(4, column, header)
            cell.font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
            cell.fill = PatternFill("solid", fgColor="8D3B53")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.comment = Comment(comment, "DM66D")
            sheet.column_dimensions[cell.column_letter].width = width
        sheet.row_dimensions[4].height = 36
        border = Border(bottom=Side(style="hair", color="E8D8DD"))
        for row in range(5, 55):
            sheet.row_dimensions[row].height = 30
            for column in range(1, last_column + 1):
                cell = sheet.cell(row, column)
                cell.number_format = "@"
                cell.font = Font(name="Calibri", size=11, color="3E3438")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border
                cell.fill = PatternFill("solid", fgColor="FFF5F7" if row % 2 else "FFFFFF")
        sheet.freeze_panes = "C5"
        sheet.auto_filter.ref = f"A4:{sheet.cell(54, last_column).column_letter}54"
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = "1:4"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        if title == "Lời chúc":
            validation = DataValidation(type="whole", operator="between", formula1=0, formula2=2147483647, allow_blank=True)
            validation.errorTitle = "Thứ tự không hợp lệ"
            validation.error = "Nhập số nguyên từ 0 trở lên."
            validation.showErrorMessage = True
            sheet.add_data_validation(validation)
            validation.add("F5:F10004")
            for row in range(5, 55):
                sheet.cell(row, 6).number_format = "0"
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)


if __name__ == "__main__":
    create_template(Path(__file__).resolve().parents[1] / "public/downloads/DM66D-wishes-template.xlsx")
    print("Đã tạo Excel mẫu trống gồm 2 sheet: Người nhận, Lời chúc.")
