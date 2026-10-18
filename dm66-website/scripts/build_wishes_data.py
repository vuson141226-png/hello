#!/usr/bin/env python3
"""Convert a checked Excel workbook into the website's simple lookup data.

The generated JSON is a public website asset. This is deliberately a local,
offline conversion: no Supabase, credentials, or uploaded source workbook is
needed at runtime. Diagnostics contain cell locations and counts only.
"""

import argparse
from collections import Counter
import json
from pathlib import Path
import re
import sys
import unicodedata
from urllib.parse import urlsplit
from zipfile import BadZipFile, ZipFile

try:
    import defusedxml  # noqa: F401: openpyxl uses hardened XML parsing
    from openpyxl import load_workbook
except ImportError:
    print("Cài scripts/requirements-import.txt trong virtualenv trước.", file=sys.stderr)
    raise SystemExit(2)


RECIPIENT_HEADERS = (
    "ID người nhận *", "Họ và tên *", "Mã sinh viên *", "Tên hiển thị",
    "Ảnh (URL hoặc đường dẫn)", "Giới thiệu", "Tính cách", "Kỷ niệm",
)
WISH_HEADERS = (
    "ID lời chúc *", "ID người nhận *", "Tiêu đề", "Người gửi", "Nội dung lời chúc *", "Thứ tự",
)
ROSTER_SHEET = "Danh sách & lời chúc"
DETAIL_SHEET = "Chi tiết 53 lời chúc"
ROSTER_HEADERS = (
    "STT", "Họ và tên", "MSV", "Người viết được phân công",
    "Người viết 1", "Lời chúc 1 (nguyên văn)", "Người viết 2", "Lời chúc 2 (nguyên văn)",
    "Người viết 3", "Lời chúc 3 (nguyên văn)", "Số lời chúc", "Còn thiếu từ", "Trạng thái",
)
DETAIL_HEADERS = (
    "STT trong danh sách", "Bạn nữ nhận lời chúc", "MSV", "Người viết",
    "Nội dung lời chúc đầy đủ", "Tệp PDF nguồn", "Trang PDF",
)
PLACEHOLDER_IMAGE = "/images/portrait-placeholder.svg"
ID_PATTERN = re.compile(r"[A-Za-z0-9_-]{1,80}\Z", re.ASCII)
CODE_PATTERN = re.compile(r"[A-Za-z0-9._-]{3,64}\Z", re.ASCII)


class WorkbookProblem(ValueError):
    """A diagnostic without the contents of a workbook cell."""


def location(cell) -> str:
    return f"{cell.parent.title}!{cell.coordinate}"


def normalize_name(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", value)
    unaccented = "".join(character for character in decomposed if not "\u0300" <= character <= "\u036f")
    return " ".join(unaccented.replace("đ", "d").replace("Đ", "d").lower().split())


def text(cell, *, required=False, preserve=False, maximum=10000) -> str:
    if cell.data_type in {"f", "e"}:
        raise WorkbookProblem(f"{location(cell)}: nhập văn bản, không dùng công thức hoặc ô lỗi.")
    if cell.value is None:
        value = ""
    elif not isinstance(cell.value, str):
        raise WorkbookProblem(f"{location(cell)}: nhập dạng Văn bản (Text), không dùng số hoặc ngày tháng.")
    else:
        value = cell.value if preserve else cell.value.strip()
    if required and not value.strip():
        raise WorkbookProblem(f"{location(cell)}: thiếu nội dung bắt buộc.")
    if len(value) > maximum or any(ord(char) < 32 and char not in "\n\r\t" for char in value):
        raise WorkbookProblem(f"{location(cell)}: văn bản quá dài hoặc có ký tự điều khiển không hợp lệ.")
    return value


def student_code(cell) -> str:
    value = text(cell, maximum=64)
    if value and not CODE_PATTERN.fullmatch(value):
        raise WorkbookProblem(f"{location(cell)}: mã sinh viên phải là văn bản gồm chữ, số, chấm, gạch ngang hoặc gạch dưới.")
    return value


def identifier(cell) -> str:
    value = text(cell, required=True, maximum=80)
    if not ID_PATTERN.fullmatch(value):
        raise WorkbookProblem(f"{location(cell)}: ID chỉ gồm chữ ASCII, số, gạch ngang hoặc gạch dưới.")
    return value


def positive_number(cell) -> int:
    value = cell.value
    if cell.data_type in {"f", "e"} or type(value) is not int or not 1 <= value <= 1000000:
        raise WorkbookProblem(f"{location(cell)}: STT phải là số nguyên dương.")
    return value


def rows(sheet, headers, header_row, maximum):
    if sheet.max_column != len(headers) or sheet.max_row > maximum + header_row:
        raise WorkbookProblem(f"{sheet.title}: số dòng/cột không phù hợp với mẫu.")
    actual = tuple(cell.value for cell in sheet[header_row])
    if actual != headers:
        raise WorkbookProblem(f"{sheet.title}: giữ nguyên tiêu đề và thứ tự cột ở dòng {header_row}.")
    for cells in sheet.iter_rows(min_row=header_row + 1):
        if any(cell.value is not None and cell.value != "" for cell in cells):
            yield cells


def new_recipient(recipient_id, full_name, code):
    if not normalize_name(full_name):
        raise WorkbookProblem("Họ và tên không được trống sau chuẩn hóa.")
    return {
        "id": recipient_id,
        "fullName": full_name,
        "displayName": full_name,
        "studentCode": code,
        "image": PLACEHOLDER_IMAGE,
        "introduction": "",
        "personality": "",
        "wishes": [],
    }


def check_unique_recipients(recipients):
    seen_ids, seen_pairs = set(), set()
    for recipient in recipients:
        pair = (normalize_name(recipient["fullName"]), recipient["studentCode"].upper())
        if recipient["id"] in seen_ids or pair in seen_pairs:
            raise WorkbookProblem("ID hoặc cặp họ tên/mã sinh viên người nhận bị lặp.")
        seen_ids.add(recipient["id"])
        seen_pairs.add(pair)


def read_original(workbook):
    recipients, roster = [], {}
    roster_messages = Counter()
    for cells in rows(workbook[ROSTER_SHEET], ROSTER_HEADERS, 5, 1000):
        number = positive_number(cells[0])
        if number in roster:
            raise WorkbookProblem(f"{location(cells[0])}: STT bị lặp.")
        recipient = new_recipient(
            f"recipient-{number:03d}", text(cells[1], required=True, maximum=120), student_code(cells[2]),
        )
        recipients.append(recipient)
        roster[number] = recipient
        for sender_column in (4, 6, 8):
            sender = text(cells[sender_column], maximum=120)
            message = text(cells[sender_column + 1], preserve=True)
            if message.strip():
                if not sender:
                    raise WorkbookProblem(f"{location(cells[sender_column])}: lời chúc có nội dung nhưng thiếu người viết.")
                roster_messages[(number, sender, message)] += 1
            elif sender:
                raise WorkbookProblem(f"{location(cells[sender_column + 1])}: có người viết nhưng thiếu nội dung lời chúc.")
    check_unique_recipients(recipients)
    detail_messages = Counter()
    blank_detail_codes = 0
    for number, cells in enumerate(rows(workbook[DETAIL_SHEET], DETAIL_HEADERS, 4, 10000), 1):
        roster_number = positive_number(cells[0])
        recipient = roster.get(roster_number)
        if recipient is None:
            raise WorkbookProblem(f"{location(cells[0])}: STT chưa có trong danh sách người nhận.")
        full_name = text(cells[1], required=True, maximum=120)
        code = student_code(cells[2])
        if normalize_name(full_name) != normalize_name(recipient["fullName"]):
            raise WorkbookProblem(f"{location(cells[1])}: họ tên không khớp người nhận có cùng STT.")
        if code and code.upper() != recipient["studentCode"].upper():
            raise WorkbookProblem(f"{location(cells[2])}: mã sinh viên không khớp danh sách người nhận.")
        if not code:
            blank_detail_codes += 1
        sender = text(cells[3], required=True, maximum=120)
        message = text(cells[4], required=True, preserve=True)
        detail_messages[(roster_number, sender, message)] += 1
        recipient["wishes"].append({
            "id": f"wish-{number:03d}", "title": f"Lời chúc từ {sender}",
            "senderName": sender, "message": message,
        })
    if roster_messages != detail_messages:
        raise WorkbookProblem("Nội dung/người viết hoặc số lượng lời chúc giữa danh sách và sheet chi tiết không khớp; cần kiểm tra file nguồn.")
    return recipients, {"blankDetailCodes": blank_detail_codes, "sourceFormat": "original"}


def read_template(workbook):
    recipients, by_id = [], {}
    for cells in rows(workbook["Người nhận"], RECIPIENT_HEADERS, 4, 1000):
        recipient = new_recipient(identifier(cells[0]), text(cells[1], required=True, maximum=120), student_code(cells[2]))
        recipient["displayName"] = text(cells[3], maximum=120) or recipient["fullName"]
        image = text(cells[4], maximum=500)
        parsed = urlsplit(image)
        if image and not (
            (image.startswith("/") and not image.startswith("//")) or
            (parsed.scheme == "https" and parsed.hostname and not parsed.username and not parsed.password)
        ):
            raise WorkbookProblem(f"{location(cells[4])}: ảnh phải là đường dẫn /… hoặc URL HTTPS.")
        recipient["image"] = image or PLACEHOLDER_IMAGE
        recipient["introduction"] = text(cells[5], preserve=True)
        recipient["personality"] = text(cells[6], preserve=True)
        memory = text(cells[7], preserve=True)
        if memory:
            recipient["memory"] = memory
        recipients.append(recipient)
        by_id[recipient["id"]] = recipient
    check_unique_recipients(recipients)
    seen_wishes = set()
    sorted_wishes = []
    for row_number, cells in enumerate(rows(workbook["Lời chúc"], WISH_HEADERS, 4, 10000)):
        wish_id = identifier(cells[0])
        if wish_id in seen_wishes:
            raise WorkbookProblem(f"{location(cells[0])}: ID lời chúc bị lặp.")
        recipient_id = identifier(cells[1])
        if recipient_id not in by_id:
            raise WorkbookProblem(f"{location(cells[1])}: ID chưa có trong sheet Người nhận.")
        sender = text(cells[3], maximum=120) or "Một người bạn DM66D"
        order = cells[5].value
        if order is None or order == "":
            order = row_number
        if isinstance(order, str) and re.fullmatch(r"[0-9]{1,10}", order.strip(), re.ASCII):
            order = int(order.strip())
        if cells[5].data_type in {"f", "e"} or type(order) is not int or not 0 <= order <= 2147483647:
            raise WorkbookProblem(f"{location(cells[5])}: thứ tự phải là số nguyên không âm.")
        wish = {
            "id": wish_id, "title": text(cells[2], maximum=200) or f"Lời chúc từ {sender}",
            "senderName": sender, "message": text(cells[4], required=True, preserve=True),
        }
        sorted_wishes.append((order, row_number, recipient_id, wish))
        seen_wishes.add(wish_id)
    for _, _, recipient_id, wish in sorted(sorted_wishes, key=lambda entry: entry[:2]):
        by_id[recipient_id]["wishes"].append(wish)
    return recipients, {"blankDetailCodes": 0, "sourceFormat": "template"}


def validate_archive(path):
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 20 * 1024 * 1024:
        raise WorkbookProblem("Cần file .xlsx không quá 20 MiB.")
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 80 * 1024 * 1024:
            raise WorkbookProblem("Workbook có cấu trúc quá lớn.")
        if any("vbaproject" in entry.filename.lower() or entry.filename.startswith("xl/externalLinks/") for entry in entries):
            raise WorkbookProblem("Không nhận macro hoặc liên kết workbook ngoài.")


def build_data(path):
    validate_archive(path)
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        if set(workbook.sheetnames) == {"Người nhận", "Lời chúc"}:
            recipients, notes = read_template(workbook)
        elif {ROSTER_SHEET, DETAIL_SHEET}.issubset(workbook.sheetnames):
            recipients, notes = read_original(workbook)
        else:
            raise WorkbookProblem("Cần mẫu hai sheet Người nhận/Lời chúc hoặc workbook tổng hợp có danh sách và sheet chi tiết.")
        distribution = Counter(len(recipient["wishes"]) for recipient in recipients)
        summary = {
            "recipients": len(recipients),
            "wishes": sum(len(recipient["wishes"]) for recipient in recipients),
            "missingStudentCodes": sum(not recipient["studentCode"] for recipient in recipients),
            "wishCountDistribution": dict(sorted(distribution.items())),
            **notes,
        }
        return recipients, summary
    finally:
        workbook.close()


def report_markdown(summary):
    lines = [
        "# Đối chiếu dữ liệu lời chúc", "",
        "Dữ liệu được tạo từ workbook người dùng cung cấp; giữ nguyên nội dung, dấu và xuống dòng của từng lời chúc.", "",
        f"- Người nhận: {summary['recipients']}.",
        f"- Lời chúc: {summary['wishes']}.",
        f"- Người nhận thiếu mã sinh viên trong danh sách chính: {summary['missingStudentCodes']}.",
        "", "| Số lời chúc mỗi người | Số người nhận |", "| --- | --- |",
    ]
    lines.extend(f"| {count} | {total} |" for count, total in summary["wishCountDistribution"].items())
    lines += ["", "Các người nhận chưa có lời chúc vẫn nằm trong danh sách để giao diện hiển thị đúng số lượng 0."]
    if summary["sourceFormat"] == "original":
        lines += [
            "",
            "Danh sách chính xác định người nhận bằng STT, họ tên và mã sinh viên. Sheet chi tiết cung cấp nguyên văn từng lời chúc; chương trình kiểm tra họ tên/mã có sẵn và đối chiếu toàn bộ người viết/nội dung với danh sách chính.",
            "",
            f"Có {summary['blankDetailCodes']} dòng chi tiết trống mã sinh viên; các dòng này được nối bằng STT và họ tên đã khớp với danh sách chính. Không tạo hay suy đoán mã mới.",
            "",
            "Ghi chú về mã sinh viên ở sheet nguồn không còn đồng nhất với các ô dữ liệu đã điền; chương trình sử dụng các ô của danh sách chính và không áp dụng ghi chú cũ để xóa mã hiện có.",
            "",
            "Workbook không cung cấp ảnh chân dung tương ứng, nên toàn bộ người nhận dùng ảnh minh họa trung tính. Không gán ảnh từ danh sách cũ.",
        ]
    lines += [
        "",
        "JSON tạo ra ở `public/data/recipient-wishes.json` là dữ liệu công khai phục vụ tra cứu trong trình duyệt. Workbook nguồn không được sao chép vào repository. Báo cáo này chỉ chứa số lượng, không chứa mã sinh viên hay nội dung lời chúc.", "",
    ]
    return "\n".join(lines)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Đối chiếu Excel và tạo JSON công khai cho tra cứu họ tên + mã sinh viên.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("public/data/recipient-wishes.json"))
    parser.add_argument("--report", type=Path, default=Path("docs/wishes-data-report.md"))
    arguments = parser.parse_args(argv)
    try:
        recipients, summary = build_data(arguments.workbook.resolve(strict=True))
        arguments.output.parent.mkdir(parents=True, exist_ok=True)
        arguments.output.write_text(json.dumps(recipients, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        arguments.report.parent.mkdir(parents=True, exist_ok=True)
        arguments.report.write_text(report_markdown(summary), encoding="utf-8")
        print(f"Đã tạo dữ liệu: {summary['recipients']} người nhận, {summary['wishes']} lời chúc; thiếu mã: {summary['missingStudentCodes']}.")
        return 0
    except (WorkbookProblem, OSError, BadZipFile, ValueError):
        error = sys.exc_info()[1]
        # Validation errors are deliberately phrased with locations/counts only.
        print(f"Lỗi: {error}" if isinstance(error, WorkbookProblem) else "Không đọc/ghi được workbook hoặc tệp đầu ra.", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
