"""Security and format checks using generated, entirely fictitious fixtures."""

from contextlib import redirect_stderr, redirect_stdout
import io
import json
from pathlib import Path
import tempfile
import unittest
from urllib.error import HTTPError

from openpyxl import load_workbook

from create_private_wishes_template import create_template
from import_private_wishes import (
    ImportProblem, NoRedirect, PRIVATE_DIRECTORY, REPOSITORY, apply_import,
    ensure_private_location, main, normalize_name, validate_workbook,
)


class FakeResponse:
    def __init__(self, payload):
        self.payload = payload

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def read(self, _maximum):
        return json.dumps(self.payload).encode()


class FakeOpener:
    def __init__(self, error=None):
        self.request = None
        self.error = error

    def open(self, request, timeout):
        self.request = request
        if self.error:
            raise self.error
        payload = json.loads(request.data)
        return FakeResponse({"recipients": len(payload["p_recipients"]), "wishes": len(payload["p_wishes"])})


class PrivateWishesImportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "filled.private.xlsx"
        create_template(self.path)

    def fixture(self):
        workbook = load_workbook(self.path)
        recipients = workbook["Người nhận"]
        recipients.append([])
        for column, value in enumerate([
            "synthetic-recipient", "Người thử A", "DEMO-NOT-A-STUDENT-001", "Tên thử",
            "/images/placeholder.svg", "", "", "",
        ], 1):
            recipients.cell(5, column, value)
        for row in [5, 6]:
            for column, value in enumerate([
                f"synthetic-wish-{row}", "synthetic-recipient", f"Lời chúc thử {row}", "Người gửi thử",
                f"Nội dung giả lập cho kiểm thử {row}.", row - 5,
            ], 1):
                workbook["Lời chúc"].cell(row, column, value)
        workbook.save(self.path)
        workbook.close()

    def change(self, sheet, coordinate, value):
        workbook = load_workbook(self.path)
        workbook[sheet][coordinate] = value
        workbook.save(self.path)
        workbook.close()

    def test_blank_public_template_has_only_two_exact_sheets_and_no_records(self):
        self.assertEqual(validate_workbook(self.path), ([], []))
        workbook = load_workbook(self.path)
        self.assertEqual(workbook.sheetnames, ["Người nhận", "Lời chúc"])
        self.assertEqual(workbook["Người nhận"]["C5"].number_format, "@")
        workbook.close()

    def test_many_wishes_share_stable_recipient_id_and_name_normalization(self):
        self.fixture()
        recipients, wishes = validate_workbook(self.path)
        self.assertEqual((len(recipients), len(wishes)), (1, 2))
        self.assertTrue(all(wish["recipient_id"] == recipients[0]["id"] for wish in wishes))
        self.assertEqual(normalize_name("  ĐẶNG\t Thị   Ánh  "), "dang thi anh")

    def test_numeric_student_codes_rejected_before_losing_leading_zero(self):
        self.fixture()
        self.change("Người nhận", "C5", 12345)
        with self.assertRaisesRegex(ImportProblem, r"Người nhận!C5"):
            validate_workbook(self.path)

    def test_name_fields_reject_internal_controls_before_database_write(self):
        self.fixture()
        for coordinate, value in [("B5", "Người\nThử"), ("B5", "Người\tThử"), ("D5", "Tên\x7fThử")]:
            with self.subTest(coordinate=coordinate):
                self.change("Người nhận", coordinate, value)
                with self.assertRaisesRegex(ImportProblem, r"Người nhận!"):
                    validate_workbook(self.path)
                self.change("Người nhận", coordinate, "Tên thử")

    def test_incomplete_row_has_safe_error_instead_of_read_only_cell_crash(self):
        self.fixture()
        self.change("Người nhận", "C5", None)
        with self.assertRaisesRegex(ImportProblem, r"Người nhận!C5"):
            validate_workbook(self.path)

    def test_formula_cells_are_not_evaluated_or_imported(self):
        self.fixture()
        self.change("Lời chúc", "E5", '=HYPERLINK("https://example.invalid")')
        with self.assertRaisesRegex(ImportProblem, r"Lời chúc!E5"):
            validate_workbook(self.path)

    def test_duplicate_wish_ids_are_rejected(self):
        self.fixture()
        self.change("Lời chúc", "A6", "synthetic-wish-5")
        with self.assertRaisesRegex(ImportProblem, r"Lời chúc!A6"):
            validate_workbook(self.path)

    def test_duplicate_login_pairs_are_rejected_even_with_different_ids(self):
        self.fixture()
        workbook = load_workbook(self.path)
        for column, cell in enumerate(workbook["Người nhận"][5], 1):
            workbook["Người nhận"].cell(6, column, cell.value)
        workbook["Người nhận"]["A6"] = "another-synthetic-recipient"
        workbook["Người nhận"]["B6"] = "nguoi thu a"
        workbook.save(self.path)
        workbook.close()
        with self.assertRaisesRegex(ImportProblem, r"Người nhận!C6"):
            validate_workbook(self.path)

    def test_unknown_recipient_is_rejected(self):
        self.fixture()
        self.change("Lời chúc", "B6", "missing-synthetic-recipient")
        with self.assertRaisesRegex(ImportProblem, r"Lời chúc!B6"):
            validate_workbook(self.path)

    def test_sort_order_and_image_scheme_are_checked(self):
        self.fixture()
        self.change("Lời chúc", "F5", -1)
        with self.assertRaisesRegex(ImportProblem, r"Lời chúc!F5"):
            validate_workbook(self.path)
        self.change("Lời chúc", "F5", 0)
        self.change("Người nhận", "E5", "javascript:alert(1)")
        with self.assertRaisesRegex(ImportProblem, r"Người nhận!E5"):
            validate_workbook(self.path)

    def test_dry_run_never_opens_network_or_prints_private_contents(self):
        self.fixture()
        output = io.StringIO()
        with redirect_stdout(output):
            self.assertEqual(main([str(self.path)]), 0)
        self.assertIn("1 người nhận, 2 lời chúc", output.getvalue())
        self.assertNotIn("DEMO-NOT-A-STUDENT", output.getvalue())
        self.assertNotIn("Nội dung giả lập", output.getvalue())

    def test_apply_uses_one_atomic_rpc_and_legacy_bearer_header(self):
        self.fixture()
        recipients, wishes = validate_workbook(self.path)
        opener = FakeOpener()
        apply_import(recipients, wishes, {"SUPABASE_URL": "https://synthetic.supabase.co", "SUPABASE_SERVICE_ROLE_KEY": "proxy-placeholder", "SUPABASE_KEY_FORMAT": "legacy"}, opener)
        self.assertTrue(opener.request.full_url.endswith("/rpc/dm66_import_workbook"))
        self.assertEqual(opener.request.get_header("Authorization"), "Bearer proxy-placeholder")
        self.assertEqual(set(json.loads(opener.request.data)), {"p_recipients", "p_wishes"})

    def test_secret_key_format_does_not_use_bearer(self):
        self.fixture()
        opener = FakeOpener()
        apply_import(*validate_workbook(self.path), {"SUPABASE_URL": "https://synthetic.supabase.co", "SUPABASE_SERVICE_ROLE_KEY": "sb_secret_synthetic"}, opener)
        self.assertIsNone(opener.request.get_header("Authorization"))

    def test_server_error_body_cannot_leak_private_content(self):
        self.fixture()
        error = HTTPError("https://synthetic.supabase.co", 400, "PRIVATE-MESSAGE", {}, io.BytesIO(b"PRIVATE-CODE"))
        with self.assertRaises(ImportProblem) as raised:
            apply_import(*validate_workbook(self.path), {"SUPABASE_URL": "https://synthetic.supabase.co", "SUPABASE_SERVICE_ROLE_KEY": "synthetic-key"}, FakeOpener(error))
        self.assertNotIn("PRIVATE", str(raised.exception))

    def test_plain_http_and_non_supabase_destinations_are_rejected(self):
        self.fixture()
        for url in ["http://synthetic.supabase.co", "https://supabase.co.evil.invalid", "https://user:pass@synthetic.supabase.co", "https://synthetic.supabase.co/path"]:
            with self.subTest(url=url), self.assertRaises(ImportProblem):
                apply_import(*validate_workbook(self.path), {"SUPABASE_URL": url, "SUPABASE_SERVICE_ROLE_KEY": "synthetic-key"}, FakeOpener())

    def test_redirects_never_forward_private_body_or_credentials(self):
        self.assertIsNone(NoRedirect().redirect_request(None, None, 307, "", {}, "https://elsewhere.invalid"))

    def test_filled_files_refused_in_public_repo_locations(self):
        with self.assertRaises(ImportProblem):
            ensure_private_location(REPOSITORY / "dm66-website/public/filled.xlsx")
        ensure_private_location(PRIVATE_DIRECTORY / "filled.xlsx")

    def test_missing_input_has_no_traceback_or_path_echo(self):
        output = io.StringIO()
        with redirect_stderr(output):
            self.assertEqual(main([str(Path(self.directory.name) / "secret-person.xlsx")]), 1)
        self.assertNotIn("secret-person", output.getvalue())
        self.assertNotIn("Traceback", output.getvalue())


if __name__ == "__main__":
    unittest.main()
